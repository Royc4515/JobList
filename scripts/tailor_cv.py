#!/usr/bin/env python3
"""Build Roy's tailored CV versions (one per entry in SUMMARIES) from his base CV (.docx).

Only the summary, the graduation date, and the order of a few lines change;
everything else is copied from the base file as-is. Roy approved these exact
versions on 2026-09-30 (see submissions/2026-09-30-packages.md).

Usage:
    pip install python-docx
    python3 scripts/tailor_cv.py <path to Roy_Carmelli_CV_sep2026.docx> <output dir>

Then export each .docx to PDF (Word or LibreOffice) and check it is one page.
The output contains personal details - never commit it (this repo is public).
"""
import sys
from pathlib import Path

import docx

B = True
SUMMARIES = {
    "Google": [
        ("Third-year ", None), ("Computer Science & Neuroscience", B),
        (" B.Sc. student at Bar-Ilan University, looking for a", None),
        (" software engineering internship", B), (", especially in", None),
        (" AI systems", B), (" and ", None), ("backend", B), ("", None), ("", None),
        (". I build and ship real projects on my own, from web apps to AI tools, "
         "and I am available 2-3 days a week. ", None),
    ],
    "Marvell": [
        ("Third-year ", None), ("Computer Science & Neuroscience", B),
        (" B.Sc. student at Bar-Ilan University, looking for a", None),
        (" student ", B), ("role in", None), (" data and AI infrastructure", B),
        (". I enjoy turning ", None), ("messy real-world data", B), (" into clean, ", None),
        ("model-ready datasets", B),
        (", and I build and ship real projects on my own, from data pipelines to AI tools. ", None),
    ],
    "FIL": [
        ("Third-year ", None), ("Computer Science & Neuroscience", B),
        (" B.Sc. student at Bar-Ilan University, looking for a", None),
        (" student ", B), ("role in", None), (" backend development", B), (": ", None),
        ("APIs and services", B), (" for ", None), ("real systems", B),
        (". I build and ship real projects on my own, and I pick up whatever stack "
         "or tool the job requires. ", None),
    ],
    "OptimumEDA": [
        ("Third-year ", None), ("Computer Science & Neuroscience", B),
        (" B.Sc. student at Bar-Ilan University, looking for a", None),
        (" student ", B), ("role in", None), (" software and automation", B), (". ", None),
        ("Python", B), (" is my main language, and much of what I build is ", None),
        ("automation", B),
        (": scripts and tools that take repetitive work off people's hands. ", None),
    ],
    "AmazonMLIL": [
        ("Third-year ", None), ("Computer Science & Neuroscience", B),
        (" B.Sc. student at Bar-Ilan University, looking for a", None),
        (" student ", B), ("role in", None), (" ML infrastructure", B), (" and ", None),
        ("automation", B), (". I write ", None), ("Python", B),
        (" daily and build tools with AI coding agents, and I pick up whatever stack "
         "or tool the job requires. ", None),
    ],
    "Claroty": [
        ("Third-year ", None), ("Computer Science & Neuroscience", B),
        (" B.Sc. student at Bar-Ilan University, looking for a", None),
        (" student ", B), ("role in", None), (" software engineering", B), (", ", None),
        ("Python", B), (" and ", None), ("AI tooling", B), (" first. I build agents and tools "
         "that have to work end to end, and I am available 2-3 days a week. ", None),
    ],
    "F5": [
        ("Third-year ", None), ("Computer Science & Neuroscience", B),
        (" B.Sc. student at Bar-Ilan University, looking for a", None),
        (" software development internship", B), (", ", None),
        ("Python", B), (" and ", None), ("backend/web", B),
        (" first, into ", None), ("web application security", B),
        (". I build tools that call many APIs (Aside, a browser extension over six "
         "model providers), use Git daily, and have about 20 hours a week. ", None),
    ],
}


def find(paragraphs, needle):
    for i, p in enumerate(paragraphs):
        if needle in p.text:
            return i
    raise SystemExit(f"Base CV changed: could not find {needle!r}")


def build(base, name, runs):
    d = docx.Document(base)
    P = d.paragraphs
    summary = P[find(P, "looking for a")].runs
    for r, (text, bold) in zip(summary, runs + [("", None)] * (len(summary) - len(runs))):
        r.text, r.bold = text, bold
    edu = P[find(P, "Oct 2024 - present")]
    for r in edu.runs:
        r.text = r.text.replace("expected 2027", "expected Feb 2028")
    if name in ("Marvell", "AmazonMLIL"):  # Python Data Processing first in coursework
        c = P[find(P, "Relevant coursework")].runs
        c[1].text, c[3].text = c[3].text, c[1].text
        c[2].text, c[4].text = c[4].text, c[2].text
    if name == "FIL":  # backend bullet first in the Wolt project
        P = d.paragraphs
        P[find(P, "Owned the")]._p.addprevious(P[find(P, "Contributed across the")]._p)
    return d


def main():
    if len(sys.argv) != 3:
        raise SystemExit(__doc__)
    base, out = Path(sys.argv[1]), Path(sys.argv[2])
    out.mkdir(parents=True, exist_ok=True)
    for name, runs in SUMMARIES.items():
        path = out / f"Roy_Carmelli_CV_{name}.docx"
        build(base, name, runs).save(path)
        print("wrote", path)


if __name__ == "__main__":
    main()
